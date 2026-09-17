"""LCA Inventory Assessment Excel generator.

Public function
---------------
export_lca_inventory(system, route, economics) -> pathlib.Path

Generates `lca/{RouteTitle}_InventoryAssessment.xlsx` with two sheets:
  "Capital Goods"      — NAICS-allocated purchase costs (2012 USD/kg SCP)
  "Operational Flows"  — stream/utility flows from BioSTEAM per row spec

All values computed directly from BioSTEAM stream/unit objects after
system.simulate() — no external links, no hardcoded numbers.

Framework §8 / Plan: LCA Automation.
"""

from __future__ import annotations

import pathlib

import biosteam as bst
import openpyxl
from openpyxl.styles import Alignment, Font

from common.lca_config import (
    COST_KEY_TO_NAICS_ROW,
    ELEMENTARY_FLOWS,
    NAICS_CATEGORIES,
    OP_FLOW_SPECS,
    _CEPCI_LCA_BASE,
)
from common.nutrients import build_nutrients_properties
from common.operating_hours import effective_operating_hours
from common.parameters import (
    ECONOMICS,
    NUTRIENTS_RECIPE,
    PRODUCTION_TARGET_MT_YR,
    EconomicBasis,
)

# ---------------------------------------------------------------------------
# Internal constants
# ---------------------------------------------------------------------------

_MW_H2O: float = 18.015       # kg/kmol — steam / chilled-water molar mass
_RHO_H2O: float = 1000.0      # kg/m³  — liquid water density for m³ conversion

_ROUTE_TITLES: dict[str, str] = {
    'fructose':         'Fructose',
    'acetate':          'Acetate',
    'formate':          'Formate',
    'gas_fermentation': 'GasFerment',
}


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def export_lca_inventory(
    system: bst.System,
    route: str,
    economics: EconomicBasis,
) -> pathlib.Path:
    """Generate the LCA Inventory Assessment workbook for one route.

    Writes `lca/{RouteTitle}_InventoryAssessment.xlsx` (self-contained, no
    external links).  Returns the output path.

    Parameters
    ----------
    system : bst.System
        Fully simulated BioSTEAM system for this route.
    route : str
        Route key ('fructose', 'acetate', 'formate', or 'gas_fermentation').
    economics : EconomicBasis
        Shared economic parameters (CEPCI, plant life, etc.).

    Returns
    -------
    pathlib.Path
        Absolute path to the written workbook.
    """
    if route not in _ROUTE_TITLES:
        raise ValueError(
            f"Unknown route {route!r}. Expected one of {list(_ROUTE_TITLES)}."
        )
    if route not in OP_FLOW_SPECS:
        raise KeyError(
            f"No OP_FLOW_SPECS defined for route {route!r}."
        )

    # ------------------------------------------------------------------
    # Runtime derived quantities
    # ------------------------------------------------------------------
    # composite_price: mass-weighted $/kg nutrients, computed from NUTRIENTS_RECIPE.
    # Same call pattern used by all four model builders (common/nutrients.py).
    # compute, don't assert — Framework §hard-rules.
    composite_price, _, nutrients_mass_fracs = build_nutrients_properties(
        concentrations=NUTRIENTS_RECIPE.concentrations,
        prices=NUTRIENTS_RECIPE.prices,
        achieved_titer=NUTRIENTS_RECIPE.achieved_titer,
    )

    cepci_ratio: float = _CEPCI_LCA_BASE / economics.CEPCI
    plant_life_yr: int = economics.duration[1] - economics.duration[0]
    annual_production_kg: float = PRODUCTION_TARGET_MT_YR * 1000.0   # kg/yr

    # Build fast-lookup dicts from system
    unit_map: dict[str, bst.Unit] = {u.ID: u for u in system.units}
    stream_map: dict[str, bst.Stream] = {
        s.ID: s for s in system.streams if s.ID
    }

    # ------------------------------------------------------------------
    # Capital Goods computation
    # ------------------------------------------------------------------
    naics_raw, naics_totals, naics_contents = _compute_capital_goods(
        system, economics, cepci_ratio, plant_life_yr, annual_production_kg,
    )

    # ------------------------------------------------------------------
    # Operational Flows computation
    # ------------------------------------------------------------------
    op_rows = _compute_op_flows(
        route, unit_map, stream_map, economics,
        composite_price, nutrients_mass_fracs, cepci_ratio,
    )

    # ------------------------------------------------------------------
    # SCP production rate — CNecatorBiomass from ST104 outlet (kg CDW / hr)
    # Used to normalise each Operational Flow row per kg SCP produced.
    # Same product identification as run_models.py: highest CNecatorBiomass
    # flow among system.products (excludes WWT/seed waste streams).
    # ------------------------------------------------------------------
    scp_product = max(
        (s for s in system.products if s.imass['CNecatorBiomass'] > 0.0),
        key=lambda s: s.imass['CNecatorBiomass'],
    )
    scp_kgh: float = scp_product.imass['CNecatorBiomass']   # kg CDW / hr

    # ------------------------------------------------------------------
    # Write workbook
    # ------------------------------------------------------------------
    lca_dir = pathlib.Path(__file__).parent.parent / 'lca'
    lca_dir.mkdir(exist_ok=True)
    out_path = lca_dir / f'{_ROUTE_TITLES[route]}_InventoryAssessment.xlsx'

    _write_workbook(out_path, naics_raw, naics_totals, naics_contents, op_rows, scp_kgh)
    return out_path


# ---------------------------------------------------------------------------
# Capital Goods
# ---------------------------------------------------------------------------

def _compute_capital_goods(
    system: bst.System,
    economics: EconomicBasis,
    cepci_ratio: float,
    plant_life_yr: int,
    annual_production_kg: float,
) -> tuple[list[float], list[float], list[str]]:
    """Sum purchase_costs into NAICS rows, deflate to 2012 USD, and normalise.

    Returns a tuple (raw_2025, normalized, contents):
        raw_2025   : list of six floats — total purchase cost per NAICS row in
                     the project dollar year (2025 USD), before deflation.
        normalized : list of six floats — 2012 USD per kg SCP over the plant
                     lifetime:
                         value = purchase_cost × (CEPCI_2012 / CEPCI_2025)
                                 / (plant_life_yr × annual_production_kg/yr)
        contents   : list of six strings — one line per unit, format
                     "UnitID (cost_key1, cost_key2, ...)", joined with newlines.
                     Reflects the units and cost items that map to each row.

    Raises KeyError if any purchase_costs key is not in COST_KEY_TO_NAICS_ROW,
    forcing the mapping to stay current with BioSTEAM cost item changes.
    """
    n_rows = len(NAICS_CATEGORIES)
    raw_totals:   list[float]                  = [0.0] * n_rows
    # row_membership[i] = {unit_id: [cost_key, ...]} — insertion-ordered
    row_membership: list[dict[str, list[str]]] = [{} for _ in range(n_rows)]

    for unit in system.units:
        for key, cost in unit.purchase_costs.items():
            if key not in COST_KEY_TO_NAICS_ROW:
                raise KeyError(
                    f"Purchase cost key {key!r} from unit {unit.ID!r} "
                    f"({type(unit).__name__}) is not in COST_KEY_TO_NAICS_ROW. "
                    "Add it to common/lca_config.py before running."
                )
            row_idx = COST_KEY_TO_NAICS_ROW[key]
            raw_totals[row_idx] += cost   # 2025 USD
            membership = row_membership[row_idx]
            if unit.ID not in membership:
                membership[unit.ID] = []
            membership[unit.ID].append(key)

    # Format membership dicts as human-readable strings for the xlsx column.
    # One line per unit: "UnitID (cost_key1, cost_key2)"
    contents: list[str] = [
        '\n'.join(
            f"{uid} ({', '.join(keys)})"
            for uid, keys in membership.items()
        )
        for membership in row_membership
    ]

    # Deflate to 2012 USD and normalise per kg SCP.
    # Dividing by (plant_life_yr × annual_production_kg/yr) gives the capital
    # cost attributable to producing 1 kg of SCP over the plant lifetime.
    normaliser = plant_life_yr * annual_production_kg
    normalized = [v * cepci_ratio / normaliser for v in raw_totals]
    return raw_totals, normalized, contents


# ---------------------------------------------------------------------------
# Operational Flows accessor interpreter
# ---------------------------------------------------------------------------

def _compute_op_flows(
    route: str,
    unit_map: dict[str, bst.Unit],
    stream_map: dict[str, bst.Stream],
    economics: EconomicBasis,
    composite_price: float,
    nutrients_mass_fracs: dict[str, float],
    cepci_ratio: float,
) -> list[dict]:
    """Evaluate every row spec in OP_FLOW_SPECS[route] and return result dicts.

    Each result dict has:
        section    : 'ins' or 'outs'
        material   : Col A label
        scope_str  : Col B — ', '.join(spec['scope'])
        ef_label   : Col C — elementary flow string from ELEMENTARY_FLOWS
        units      : Col D label
        value      : Col E computed value (in the stated units, per hour)
    """
    results = []
    for spec in OP_FLOW_SPECS[route]:
        raw = _evaluate_accessor(
            spec['accessor'], unit_map, stream_map, economics,
            composite_price, nutrients_mass_fracs,
        )
        value = _apply_conversion(raw, spec['accessor'][0], spec['units'], cepci_ratio)
        ef_key = spec['ef_key']
        if ef_key not in ELEMENTARY_FLOWS:
            raise KeyError(
                f"Elementary flow key {ef_key!r} not in ELEMENTARY_FLOWS. "
                "Add it to common/lca_config.py."
            )
        results.append({
            'section':   spec['section'],
            'material':  spec['material'],
            'scope_str': ', '.join(spec['scope']),
            'ef_label':  ELEMENTARY_FLOWS[ef_key],
            'units':     spec['units'],
            'value':     value,
        })
    return results


def _evaluate_accessor(
    accessor: tuple,
    unit_map: dict[str, bst.Unit],
    stream_map: dict[str, bst.Stream],
    economics: EconomicBasis,
    composite_price: float,
    nutrients_mass_fracs: dict[str, float],
) -> float:
    """Dispatch to the appropriate accessor handler.  Returns a raw numeric value
    in the accessor's natural units (kg/hr, kW, kmol/hr, or $/hr) before any
    unit conversion.
    """
    acc_type, *args = accessor

    if acc_type == 'feed_stream':
        stream_id, component = args
        return stream_map[stream_id].imass[component]   # kg/hr

    elif acc_type == 'unit_in_sum':
        tuples = args[0]   # [(unit_id, in_idx, component), ...]
        return sum(
            unit_map[uid].ins[idx].imass[comp]
            for uid, idx, comp in tuples
        )  # kg/hr

    elif acc_type == 'cost_sum':
        tuples, price_arg = args
        price = composite_price if price_arg == '<composite_price>' else price_arg
        mass_kgh = sum(
            unit_map[uid].ins[idx].imass[comp]
            for uid, idx, comp in tuples
        )
        return mass_kgh * price   # $/hr (current-year)

    elif acc_type == 'power_sum':
        unit_ids = args[0]
        return sum(
            unit_map[uid].power_utility.consumption for uid in unit_ids
        )  # kW

    elif acc_type == 'h2o_feed_sum':
        stream_ids = args[0]
        return sum(stream_map[sid].imass['H2O'] for sid in stream_ids)  # kg/hr

    elif acc_type == 'utility_sum':
        unit_ids, agent_id = args
        # Match BioSTEAM agent IDs case-insensitively with space→underscore
        # (e.g. 'low_pressure_steam' or 'chilled_water' from BioSTEAM).
        agent_norm = agent_id.lower().replace(' ', '_')
        total = 0.0
        for uid in unit_ids:
            for hu in unit_map[uid].heat_utilities:
                if hu.agent.ID.lower() == agent_norm:
                    total += hu.flow   # kmol/hr per vessel, summed across all
        return total   # kmol/hr

    elif acc_type == 'air_sum':
        unit_ids, component = args
        total = 0.0
        for uid in unit_ids:
            u = unit_map[uid]
            if hasattr(u, 'air'):
                total += u.air.imass[component]
        return total   # kg/hr

    elif acc_type == 'wwt_cost':
        wwt_uid = args[0]
        wwt_inlet = unit_map[wwt_uid].ins[0]
        organic_kgh = sum(
            wwt_inlet.imass[cid]
            for cid in wwt_inlet.chemicals.IDs
            if cid != 'H2O'
        )
        return organic_kgh * economics.wwt_organic_removal_cost   # $/hr (current-year)

    elif acc_type == 'discharge':
        component = args[0]
        return unit_map['RCY101'].outs[1].imass[component]   # kg/hr

    elif acc_type == 'gas_out_sum':
        unit_ids, component = args
        return sum(
            unit_map[uid].outs[0].imass[component] for uid in unit_ids
        )  # kg/hr; outs[0] is gas vent for AeratedBioreactor subclasses

    elif acc_type == 'dryer_vapor':
        component = args[0]
        return unit_map['D101'].outs[0].imass[component]   # kg/hr

    elif acc_type == 'discharge_unlump':
        nutrient_key = args[0]
        nutrients_kgh = unit_map['RCY101'].outs[1].imass['Nutrients']
        return nutrients_kgh * nutrients_mass_fracs[nutrient_key]   # kg/hr

    elif acc_type == 'unit_out':
        unit_id, out_idx, component = args
        return unit_map[unit_id].outs[out_idx].imass[component]   # kg/hr

    elif acc_type == 'unit_out_unlump':
        unit_id, out_idx, nutrient_key = args
        nutrients_kgh = unit_map[unit_id].outs[out_idx].imass['Nutrients']
        return nutrients_kgh * nutrients_mass_fracs[nutrient_key]   # kg/hr

    elif acc_type == 'unit_out_sum':
        unit_outlet_pairs, component = args
        # unit_outlet_pairs = [(unit_id, out_idx), ...]
        return sum(
            unit_map[uid].outs[idx].imass[component]
            for uid, idx in unit_outlet_pairs
        )  # kg/hr

    elif acc_type == 'multi_out_unlump':
        unit_outlet_pairs, nutrient_key = args
        # Sum Nutrients pseudo-component across multiple outlets, then unlump
        # via mass fraction.  Framework §hard-rules: compute, don't assert.
        total_nutrients = sum(
            unit_map[uid].outs[idx].imass['Nutrients']
            for uid, idx in unit_outlet_pairs
        )
        return total_nutrients * nutrients_mass_fracs[nutrient_key]  # kg/hr

    else:
        raise ValueError(
            f"Unknown accessor type {acc_type!r}. "
            "Add a handler in common/lca_export.py _evaluate_accessor()."
        )


def _apply_conversion(
    raw: float,
    acc_type: str,
    units: str,
    cepci_ratio: float,
) -> float:
    """Convert the raw accessor value to the stated units.

    Conversion rules (determined by accessor type × units field):
        power_sum           → kW ÷ 1000 = MWh/hr (= MW)
        h2o_feed_sum  + m3  → kg ÷ 1000 = m³/hr
        utility_sum   + kg  → kmol × 18.015 = kg/hr    (steam)
        utility_sum   + m3  → kmol × 18.015 / 1000 = m³/hr  (chilled water)
        cost_sum      + 2012 USD → current $/hr × cepci_ratio = 2012 $/hr
        wwt_cost      + 2012 USD → current $/hr × cepci_ratio = 2012 $/hr
        discharge     + m3  → kg ÷ 1000 = m³/hr    (water effluent, legacy RCY101-only)
        unit_out_sum  + m3  → kg ÷ 1000 = m³/hr    (water effluent: discharge + sludge H2O)
        all others          → raw value unchanged   (kg/hr)
    """
    if acc_type == 'power_sum':
        return raw / 1000.0                         # kW → MWh/hr (= MW)

    elif acc_type == 'h2o_feed_sum' and units == 'm3':
        return raw / _RHO_H2O                       # kg → m³

    elif acc_type == 'utility_sum' and units == 'kg':
        return raw * _MW_H2O                        # kmol → kg (steam)

    elif acc_type == 'utility_sum' and units == 'm3':
        return raw * _MW_H2O / _RHO_H2O             # kmol → m³ (chilled water)

    elif acc_type in ('cost_sum', 'wwt_cost') and units == '2012 USD':
        return raw * cepci_ratio                    # current USD → 2012 USD

    elif acc_type == 'discharge' and units == 'm3':
        return raw / _RHO_H2O                       # kg → m³ (water effluent, RCY101 only — legacy)

    elif acc_type == 'unit_out_sum' and units == 'm3':
        return raw / _RHO_H2O                       # kg → m³ (water effluent: discharge + sludge H2O)

    else:
        return raw                                  # kg/hr — no further conversion


# ---------------------------------------------------------------------------
# Excel writer
# ---------------------------------------------------------------------------

def _write_workbook(
    path: pathlib.Path,
    naics_raw: list[float],
    naics_totals: list[float],
    naics_contents: list[str],
    op_rows: list[dict],
    scp_kgh: float,
) -> None:
    """Write Capital Goods and Operational Flows sheets to *path*."""
    wb = openpyxl.Workbook()

    # ---- Sheet 1: Capital Goods ----
    ws_cap = wb.active
    ws_cap.title = 'Capital Goods'

    bold    = Font(bold=True)
    wrap_top = Alignment(wrap_text=True, vertical='top')

    headers_cap = [
        'NAICS Category', 'NAICS Code',
        'Equipment included', '2025 USD (total)', '2012 USD / kg SCP',
    ]
    for col, h in enumerate(headers_cap, start=1):
        cell = ws_cap.cell(row=1, column=col, value=h)
        cell.font = bold

    for i, ((name, code), contents, raw, value) in enumerate(
        zip(NAICS_CATEGORIES, naics_contents, naics_raw, naics_totals), start=2
    ):
        ws_cap.cell(row=i, column=1, value=name).alignment = wrap_top
        ws_cap.cell(row=i, column=2, value=code).alignment = wrap_top
        cell_eq = ws_cap.cell(row=i, column=3, value=contents)
        cell_eq.alignment = wrap_top
        ws_cap.cell(row=i, column=4, value=round(raw, 2)).alignment = wrap_top
        ws_cap.cell(row=i, column=5, value=round(value, 10)).alignment = wrap_top

    # ---- Sheet 2: Operational Flows ----
    ws_op = wb.create_sheet(title='Operational Flows')

    headers_op = [
        'Material', 'Scope', 'Elementary Flow', 'Units',
        'Value (per hr)', 'Value / kg SCP',
    ]
    for col, h in enumerate(headers_op, start=1):
        cell = ws_op.cell(row=1, column=col, value=h)
        cell.font = bold

    row_idx = 2
    current_section = None
    for record in op_rows:
        section = record['section']
        if section != current_section:
            # Section header row
            label = 'INPUTS' if section == 'ins' else 'OUTPUTS'
            cell = ws_op.cell(row=row_idx, column=1, value=label)
            cell.font = bold
            row_idx += 1
            current_section = section

        ws_op.cell(row=row_idx, column=1, value=record['material'])
        ws_op.cell(row=row_idx, column=2, value=record['scope_str'])
        ws_op.cell(row=row_idx, column=3, value=record['ef_label'])
        ws_op.cell(row=row_idx, column=4, value=record['units'])
        ws_op.cell(row=row_idx, column=5, value=round(record['value'], 8))
        ws_op.cell(row=row_idx, column=6, value=round(record['value'] / scp_kgh, 10))
        row_idx += 1

    wb.save(path)
